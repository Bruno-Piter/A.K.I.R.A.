"""Unit tests for extract_text. No live Neo4j required."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.arms.routines.extract_text import (
    allowed_content_type,
    allowed_filename,
    extract_text,
)


def test_txt_roundtrip(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("hello txt", encoding="utf-8")
    assert extract_text(path=path, filename="note.txt") == "hello txt"


def test_md_roundtrip(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    path.write_text("# Title\n\nbody", encoding="utf-8")
    text = extract_text(path=path, filename="note.md")
    assert "Title" in text
    assert "body" in text


def test_latin1_txt(tmp_path: Path) -> None:
    path = tmp_path / "latin.txt"
    path.write_bytes(b"caf\xe9")
    assert "caf" in extract_text(path=path, filename="latin.txt")


def test_docx_paragraph(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Hello DOCX")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "cell value"
    doc.save(path)
    text = extract_text(path=path, filename="note.docx")
    assert "Hello DOCX" in text
    assert "cell value" in text


def test_pdf_reader_path(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    blank = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(blank)
    assert extract_text(path=blank, filename="blank.pdf") == ""

    payload = _minimal_pdf_bytes("Hello PDF")
    text_pdf = tmp_path / "hello.pdf"
    text_pdf.write_bytes(payload)
    extracted = extract_text(path=text_pdf, filename="hello.pdf")
    # pypdf may or may not recover the toy stream; the path must not crash.
    assert isinstance(extracted, str)
    if extracted:
        assert "Hello PDF" in extracted


def test_unsupported_exe_raises(tmp_path: Path) -> None:
    path = tmp_path / "tool.exe"
    path.write_bytes(b"MZ\x00\x00not a document")
    with pytest.raises(ValueError, match="unsupported"):
        extract_text(path=path, filename="tool.exe")


def test_allowed_helpers() -> None:
    assert allowed_filename("report.pdf")
    assert allowed_filename("notes.markdown")
    assert not allowed_filename("photo.png")
    assert not allowed_filename("tool.exe")
    assert allowed_content_type("application/pdf")
    assert allowed_content_type("text/markdown; charset=utf-8")
    assert not allowed_content_type("image/png")
    assert not allowed_content_type("application/octet-stream")


def test_post_documents_rejects_unsupported() -> None:
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as client:
        exe = client.post(
            "/api/documents",
            files={"file": ("bad.exe", b"MZ fake", "application/octet-stream")},
        )
        png = client.post(
            "/api/documents",
            files={"file": ("photo.png", b"\x89PNG\r\n\x1a\n", "image/png")},
        )
    assert exe.status_code == 415
    assert "pdf" in str(exe.json().get("detail", "")).lower()
    assert png.status_code == 415


def _minimal_pdf_bytes(text: str) -> bytes:
    content = f"BT /F1 12 Tf 50 50 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode("ascii")
        out += obj
        out += b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)