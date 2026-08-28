"""Extract plain text from PDF, TXT, MD, and DOCX. No OCR. LangGraph-free."""

from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

ALLOWED_SUFFIXES = {".pdf", ".txt", ".md", ".markdown", ".docx"}

_SUFFIX_KIND = {
    ".pdf": "pdf",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".docx": "docx",
}

_CONTENT_TYPE_KIND = {
    "application/pdf": "pdf",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/x-markdown": "md",
    "application/markdown": "md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}

_ACCEPTED = "pdf, txt, md, docx"


def allowed_filename(name: str | None) -> bool:
    if not name:
        return False
    return Path(name).suffix.lower() in ALLOWED_SUFFIXES


def allowed_content_type(ct: str | None) -> bool:
    if not ct:
        return False
    base = ct.split(";", 1)[0].strip().lower()
    return base in _CONTENT_TYPE_KIND


def extract_text(
    path: str | Path | None = None,
    filename: str | None = None,
    content_type: str | None = None,
    file_bytes: bytes | None = None,
) -> str:
    """Return extracted plain text. Empty PDF text layer yields '' (no OCR)."""
    path_obj = Path(path) if path is not None else None
    name = filename or (path_obj.name if path_obj is not None else None)
    data = file_bytes
    if data is None:
        if path_obj is None:
            raise ValueError("extract_text requires path or file_bytes")
        data = path_obj.read_bytes()
    if isinstance(data, str):
        data = data.encode("utf-8")
    kind = _detect_kind(name, content_type, data)
    if kind == "pdf":
        return _extract_pdf(data)
    if kind == "docx":
        return _extract_docx(data)
    if kind in {"txt", "md"}:
        return _decode_text(data)
    raise ValueError(f"unsupported document type; accepted: {_ACCEPTED}")


def _detect_kind(filename: str | None, content_type: str | None, data: bytes) -> str:
    if filename:
        suffix = Path(filename).suffix.lower()
        kind = _SUFFIX_KIND.get(suffix)
        if kind:
            return kind
    if content_type:
        base = content_type.split(";", 1)[0].strip().lower()
        kind = _CONTENT_TYPE_KIND.get(base)
        if kind:
            return kind
    sniffed = _sniff(data)
    if sniffed:
        return sniffed
    raise ValueError(f"unsupported document type; accepted: {_ACCEPTED}")


def _sniff(data: bytes) -> str | None:
    if data.startswith(b"%PDF"):
        return "pdf"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(BytesIO(data)) as zf:
                names = zf.namelist()
            if "word/document.xml" in names or any(n.startswith("word/") for n in names):
                return "docx"
        except zipfile.BadZipFile:
            return None
    return None


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        raw = page.extract_text() or ""
        text = raw.strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text
        if text and text.strip():
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                if text and text.strip():
                    parts.append(text)
    return "\n".join(parts)